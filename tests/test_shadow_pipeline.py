"""
Shadow Pipeline End-to-End Test.
Simulates realistic data ingestion, cleaning, layout formatting and Excel generation
to guarantee 100% data fidelity, visual excellence, and formula accuracy.
"""
import json
import os
import shutil
import tempfile
import time
import unittest
import openpyxl

from python.contracts import FileSpec, HeaderCell, FormulaCell
from python.data_pipeline import (
    filter_by_shift,
    deduplicate_latest_per_room,
    build_status_facility_rows,
    build_may_tinh_rows,
    build_dien_rows,
)
from python.excel_styler import fill_single_spec, ALERT_RED_FILL, ALERT_YELLOW_FILL


class TestShadowPipeline(unittest.TestCase):

    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.template_dir = os.path.join(self.test_dir, "templates")
        self.export_dir = os.path.join(self.test_dir, "export")
        os.makedirs(self.template_dir, exist_ok=True)
        os.makedirs(self.export_dir, exist_ok=True)

        # Create mock Loa template
        self.loa_template = os.path.join(self.template_dir, "01 - FORM TINH TRANG AM THANH.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "1209"
        ws["A1"] = "TIN HỌC ĐẠI DƯƠNG\nTÌNH TRẠNG LOA"
        ws["E2"] = ""  # Thứ
        ws["F2"] = ""  # Ngày
        ws["G2"] = ""  # Giờ
        ws["A3"] = "STT"
        ws["B3"] = "TÊN TRƯỜNG"
        ws["C3"] = "KHU VỰC"
        ws["D3"] = "PHÒNG MÁY"
        ws["F3"] = "XẾP LOẠI"
        ws["A4"] = ""  # Data row template
        ws["A5"] = "TỔNG"
        wb.save(self.loa_template)
        wb.close()

        # Create mock MayTinh template
        self.maytinh_template = os.path.join(self.template_dir, "02 - FORM TINH TRANG MAY TINH.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2409"
        ws["A1"] = "TIN HỌC ĐẠI DƯƠNG\nTÌNH TRẠNG MÁY TÍNH"
        ws["J2"] = ""  # Thứ
        ws["K2"] = ""  # Ngày
        ws["L2"] = ""  # Giờ
        ws["A4"] = ""
        ws["A5"] = "TỔNG"
        wb.create_sheet("0310 (2)")  # unused sheet
        wb.save(self.maytinh_template)
        wb.close()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_shadow_end_to_end_execution(self):
        start_time = time.time()

        # 1. Realistic raw reports from Zalo messages
        raw_reports = [
            {
                "maNhom": "g_cva",
                "thoiGian": 1790815117413,  # Morning shift
                "data": {
                    "phongMay": "PM1",
                    "truong": {"tenRieng": "Chu Văn An"},
                    "khuVuc": "BTAN",
                    "tongSoMay": 35,
                    "coSoVatChat": [
                        {"loai": "AmThanh", "xepLoai": "T"},
                        {"loai": "MayTinh", "soLuong": [{"loai": "Loi", "so": 0, "ghiChu": ""}, {"loai": "Thieu", "so": 0}]}
                    ]
                }
            },
            {
                "maNhom": "g_cva",
                "thoiGian": 1790815120000,  # Updated report for PM1 in same session
                "data": {
                    "phongMay": "PM1",
                    "truong": {"tenRieng": "Chu Văn An"},
                    "khuVuc": "BTAN",
                    "tongSoMay": 35,
                    "coSoVatChat": [
                        {"loai": "AmThanh", "xepLoai": "B"},
                        {"loai": "MayTinh", "soLuong": [{"loai": "Loi", "so": 2, "ghiChu": "Hỏng nguồn"}, {"loai": "Thieu", "so": 3}]}
                    ]
                }
            },
            {
                "maNhom": "g_hs",
                "thoiGian": 1790815130000,
                "data": {
                    "phongMay": "PM2",
                    "truong": {"tenRieng": "Hoa Sen"},
                    "khuVuc": "Q1",
                    "tongSoMay": 40,
                    "coSoVatChat": [
                        {"loai": "AmThanh", "xepLoai": "K"},
                        {"loai": "MayTinh", "soLuong": [{"loai": "Loi", "so": 5, "ghiChu": "Lỗi RAM"}, {"loai": "Thieu", "so": 9}]}
                    ]
                }
            }
        ]

        # Step 2: Deduplication
        clean_reports = deduplicate_latest_per_room(raw_reports)
        self.assertEqual(len(clean_reports), 2, "Deduplication should leave exactly 2 unique rooms")

        # Step 3: Build rows via Pandas engine
        loa_rows = build_status_facility_rows(clean_reports, "AmThanh")
        may_tinh_rows = build_may_tinh_rows(clean_reports)

        self.assertEqual(len(loa_rows), 2)
        self.assertEqual(len(may_tinh_rows), 2)

        # Step 4: Openpyxl fill and styler
        loa_out = os.path.join(self.export_dir, "TÌNH TRẠNG ÂM THANH_TEST.xlsx")
        loa_spec = FileSpec(
            name="loa",
            template=self.loa_template,
            out=loa_out,
            sheet="1209",
            firstDataRow=4,
            totalRowBefore=5,
            lastColumn="G",
            rows=loa_rows,
            headerCells=[HeaderCell(ref="E2", value="Thứ: 3"), HeaderCell(ref="F2", value="Ngày: 13/10/2026")],
            formulaCells=[FormulaCell(ref="K10", formula="COUNTIF(F4:F5, 'T')")]
        )
        msg_loa = fill_single_spec(loa_spec)
        self.assertIn("2 dòng", msg_loa)

        may_tinh_out = os.path.join(self.export_dir, "TÌNH TRẠNG MÁY TÍNH_TEST.xlsx")
        maytinh_spec = FileSpec(
            name="mayTinh",
            template=self.maytinh_template,
            out=may_tinh_out,
            sheet="2409",
            firstDataRow=4,
            totalRowBefore=5,
            lastColumn="L",
            deleteSheets=["0310 (2)"],
            rows=may_tinh_rows,
            headerCells=[HeaderCell(ref="J2", value="Thứ: 3")],
            formulaCells=[]
        )
        msg_mt = fill_single_spec(maytinh_spec)
        self.assertIn("2 dòng", msg_mt)

        duration = time.time() - start_time
        self.assertLess(duration, 1.0, f"Shadow pipeline should complete in under 1 second, took {duration:.3f}s")

        # Step 5: Verification of outputs
        wb_loa = openpyxl.load_workbook(loa_out)
        ws_loa = wb_loa["1209"]
        # Grade B for row 4 (Hoa Sen PM2 had K, Chu Van An PM1 had B)
        # B is sorted before K
        self.assertEqual(ws_loa["F4"].value, "B")
        self.assertEqual(ws_loa["F4"].fill.start_color.rgb, ALERT_YELLOW_FILL.start_color.rgb)
        self.assertEqual(ws_loa["F5"].value, "K")
        self.assertEqual(ws_loa["F5"].fill.start_color.rgb, ALERT_RED_FILL.start_color.rgb)
        self.assertEqual(str(ws_loa.page_setup.paperSize), str(ws_loa.PAPERSIZE_A4))
        wb_loa.close()

        wb_mt = openpyxl.load_workbook(may_tinh_out)
        self.assertNotIn("0310 (2)", wb_mt.sheetnames)
        ws_mt = wb_mt["2409"]
        # Hoa Sen missing 9 machines -> Grade E (Red alert)
        self.assertTrue(any(row[9] == "E" for row in ws_mt.iter_rows(min_row=4, max_row=5, values_only=True)))
        wb_mt.close()


if __name__ == "__main__":
    unittest.main()

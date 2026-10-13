"""
Unit tests for python/export_worker.py
"""
import json
import os
import shutil
import tempfile
import unittest
import openpyxl

from python.export_worker import run_spec_mode


class TestExportWorker(unittest.TestCase):

    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.template_path = os.path.join(self.test_dir, "template.xlsx")
        self.output_path = os.path.join(self.test_dir, "out_report.xlsx")
        self.spec_path = os.path.join(self.test_dir, "spec.json")
        self.log_path = os.path.join(self.test_dir, "worker.log")

        # Create basic template
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TEST_SHEET"
        ws["A1"] = "HEADER TITLE"
        ws["A4"] = ""
        ws["A5"] = "TỔNG"
        wb.save(self.template_path)
        wb.close()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_run_spec_mode_success(self):
        spec_content = {
            "files": [
                {
                    "name": "test_report",
                    "template": self.template_path,
                    "out": self.output_path,
                    "sheet": "TEST_SHEET",
                    "firstDataRow": 4,
                    "totalRowBefore": 5,
                    "lastColumn": "D",
                    "deleteSheets": [],
                    "rows": [
                        {"A": 1, "B": "Trường A", "C": "KV1", "D": "Tốt"},
                        {"A": 2, "B": "Trường B", "C": "KV2", "D": "Bình thường"},
                    ],
                    "headerCells": [
                        {"ref": "B1", "value": "NGÀY 13/10/2026"}
                    ],
                    "formulaCells": []
                }
            ]
        }
        with open(self.spec_path, "w", encoding="utf-8") as f:
            json.dump(spec_content, f)

        exit_code = run_spec_mode(self.spec_path, self.log_path)
        self.assertEqual(exit_code, 0)
        self.assertTrue(os.path.exists(self.output_path))
        self.assertTrue(os.path.exists(self.log_path))

        with open(self.log_path, "r", encoding="utf-8") as lf:
            log_text = lf.read()
        self.assertIn("test_report: 2 dòng", log_text)

    def test_run_spec_mode_missing_template(self):
        spec_content = {
            "files": [
                {
                    "name": "fail_report",
                    "template": os.path.join(self.test_dir, "non_existent.xlsx"),
                    "out": self.output_path,
                    "sheet": "TEST_SHEET",
                    "firstDataRow": 4,
                    "totalRowBefore": 5,
                    "lastColumn": "D",
                    "rows": []
                }
            ]
        }
        with open(self.spec_path, "w", encoding="utf-8") as f:
            json.dump(spec_content, f)

        exit_code = run_spec_mode(self.spec_path, self.log_path)
        self.assertEqual(exit_code, 1)

        with open(self.log_path, "r", encoding="utf-8") as lf:
            log_text = lf.read()
        self.assertIn("ERR", log_text)


if __name__ == "__main__":
    unittest.main()

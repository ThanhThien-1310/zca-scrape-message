"""
Unit tests for python/excel_styler.py
"""
import os
import shutil
import tempfile
import unittest
import openpyxl
from openpyxl.styles import PatternFill

from python.contracts import FileSpec, HeaderCell, FormulaCell
from python.excel_styler import (
    fill_single_spec,
    apply_visual_alerts,
    setup_a4_page_printing,
    ALERT_RED_FILL,
    ALERT_YELLOW_FILL,
)


class TestExcelStyler(unittest.TestCase):

    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.template_path = os.path.join(self.test_dir, "mock_template.xlsx")
        self.output_path = os.path.join(self.test_dir, "output_report.xlsx")

        # Create a mock template
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "REPORT"
        # Row 1-3: Headers
        ws["A1"] = "TIN HỌC ĐẠI DƯƠNG\nTÌNH TRẠNG LOA"
        ws["E2"] = ""  # Thứ
        ws["F2"] = ""  # Ngày
        ws["G2"] = ""  # Giờ
        ws["A3"] = "STT"
        ws["B3"] = "TÊN TRƯỜNG"
        ws["C3"] = "KHU VỰC"
        ws["D3"] = "PHÒNG MÁY"
        ws["F3"] = "XẾP LOẠI"

        # Row 4: initial data row
        ws["A4"] = ""
        # Row 5: Total row
        ws["A5"] = "TỔNG"
        ws["F5"] = '=COUNTIF(F4:F4, "T")'

        wb.create_sheet(title="EXTRA_SHEET")
        wb.save(self.template_path)
        wb.close()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_fill_single_spec(self):
        spec = FileSpec(
            name="loa",
            template=self.template_path,
            out=self.output_path,
            sheet="REPORT",
            firstDataRow=4,
            totalRowBefore=5,
            lastColumn="G",
            deleteSheets=["EXTRA_SHEET"],
            rows=[
                {"A": 1, "B": "TH Chu Van An", "C": "BTAN", "D": "PM1", "F": "T"},
                {"A": 2, "B": "TH Chu Van An", "C": "BTAN", "D": "PM2", "F": "K"},
                {"A": 3, "B": "TH Hoa Sen", "C": "Q1", "D": "PM1", "F": "B"},
            ],
            headerCells=[
                HeaderCell(ref="E2", value="Thứ: 3"),
                HeaderCell(ref="F2", value="Ngày: 13/10/2026"),
                HeaderCell(ref="G2", value="Thời gian: 09:00"),
            ],
            formulaCells=[
                FormulaCell(ref="K10", formula="SUM(K7:K9)"),
            ],
        )

        msg = fill_single_spec(spec)
        self.assertTrue(os.path.exists(self.output_path))
        self.assertIn("3 dòng", msg)

        # Inspect generated output
        wb_out = openpyxl.load_workbook(self.output_path)
        self.assertNotIn("EXTRA_SHEET", wb_out.sheetnames)
        ws_out = wb_out["REPORT"]

        # Check headers
        self.assertEqual(ws_out["E2"].value, "Thứ: 3")
        self.assertEqual(ws_out["F2"].value, "Ngày: 13/10/2026")
        self.assertEqual(ws_out["G2"].value, "Thời gian: 09:00")

        # Check populated rows
        self.assertEqual(ws_out["A4"].value, 1)
        self.assertEqual(ws_out["B4"].value, "TH Chu Van An")
        self.assertEqual(ws_out["F4"].value, "T")

        # Row 5 was PM2 with grade K -> should have Red alert fill
        self.assertEqual(ws_out["A5"].value, 2)
        self.assertEqual(ws_out["F5"].value, "K")
        self.assertEqual(ws_out["F5"].fill.start_color.rgb, ALERT_RED_FILL.start_color.rgb)

        # Row 6 was PM1 with grade B -> should have Yellow alert fill
        self.assertEqual(ws_out["A6"].value, 3)
        self.assertEqual(ws_out["F6"].value, "B")
        self.assertEqual(ws_out["F6"].fill.start_color.rgb, ALERT_YELLOW_FILL.start_color.rgb)

        # Total row moved down from 5 to 7
        self.assertEqual(ws_out["A7"].value, "TỔNG")

        # Formula cell
        self.assertEqual(ws_out["K10"].value, "=SUM(K7:K9)")

        # Page setup
        self.assertEqual(str(ws_out.page_setup.paperSize), str(ws_out.PAPERSIZE_A4))
        self.assertEqual(ws_out.page_setup.orientation, ws_out.ORIENTATION_LANDSCAPE)
        self.assertEqual(ws_out.page_setup.fitToWidth, 1)


if __name__ == "__main__":
    unittest.main()
